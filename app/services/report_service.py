import io
from collections import Counter
from datetime import datetime
from typing import Dict, List, Optional, Tuple
import pandas as pd
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.graphics.charts.barcharts import VerticalBarChart
from reportlab.graphics.charts.legends import Legend
from reportlab.graphics.charts.piecharts import Pie
from reportlab.graphics.shapes import Drawing, String
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from sqlalchemy import select
from sqlalchemy.orm import Session, selectinload

from app.core.time import LIMA_TZ
from app.models.customer import Cliente
from app.models.customer_request import SolicitudCliente
from app.models.order import Pedido
from app.models.order_item import DetallePedido
from app.models.product import Producto
from app.services.indicator_service import IndicatorService


class ReportService:
    @staticmethod
    def _kpi_cards(kpis: List[Tuple[str, str, str]]) -> Table:
        styles = getSampleStyleSheet()
        label_style = ParagraphStyle(
            "KpiLabel", parent=styles["Normal"], fontSize=7,
            textColor=colors.HexColor("#64748b"), leading=9,
        )
        tarjetas = []
        for label, valor, color in kpis:
            valor_style = ParagraphStyle(
                "KpiValor", parent=styles["Normal"], fontSize=16,
                textColor=colors.HexColor(color), leading=18, fontName="Helvetica-Bold",
            )
            contenido = Table(
                [[Paragraph(label.upper(), label_style)], [Paragraph(str(valor), valor_style)]],
                colWidths=[125],
            )
            contenido.setStyle(TableStyle([
                ("TOPPADDING", (0, 0), (-1, -1), 1),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 1),
                ("LEFTPADDING", (0, 0), (-1, -1), 0),
            ]))
            tarjeta = Table([[contenido]], colWidths=[135])
            tarjeta.setStyle(TableStyle([
                ("BOX", (0, 0), (-1, -1), 0.75, colors.HexColor(color)),
                ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F8FAFC")),
                ("TOPPADDING", (0, 0), (-1, -1), 8),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
                ("LEFTPADDING", (0, 0), (-1, -1), 10),
            ]))
            tarjetas.append(tarjeta)
        fila = Table([tarjetas], hAlign="LEFT")
        fila.setStyle(TableStyle([
            ("LEFTPADDING", (0, 0), (-1, -1), 0),
            ("RIGHTPADDING", (0, 0), (-1, -1), 8),
            ("TOPPADDING", (0, 0), (-1, -1), 0),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
        ]))
        return fila

    @staticmethod
    def _grafico_barras(datos: Dict[str, float], titulo: str, color: str = "#1E3A8A") -> Drawing:
        d = Drawing(460, 190)
        d.add(String(0, 178, titulo, fontSize=8, fillColor=colors.HexColor("#334155")))
        bc = VerticalBarChart()
        bc.x = 40
        bc.y = 25
        bc.height = 130
        bc.width = 400
        bc.data = [list(datos.values())]
        bc.categoryAxis.categoryNames = list(datos.keys())
        bc.categoryAxis.labels.fontSize = 7
        bc.valueAxis.labels.fontSize = 7
        bc.valueAxis.valueMin = 0
        bc.bars[0].fillColor = colors.HexColor(color)
        bc.barWidth = 12
        bc.groupSpacing = 12
        d.add(bc)
        return d

    @staticmethod
    def _grafico_barras_comparativo(
        categorias: List[str], series: Dict[str, List[float]], titulo: str
    ) -> Drawing:
        paleta = ["#1E3A8A", "#94a3b8", "#ea580c"]
        d = Drawing(460, 200)
        d.add(String(0, 188, titulo, fontSize=8, fillColor=colors.HexColor("#334155")))
        bc = VerticalBarChart()
        bc.x = 40
        bc.y = 35
        bc.height = 130
        bc.width = 340
        bc.data = list(series.values())
        bc.categoryAxis.categoryNames = categorias
        bc.categoryAxis.labels.fontSize = 7
        bc.valueAxis.labels.fontSize = 7
        bc.valueAxis.valueMin = 0
        bc.barWidth = 9
        bc.groupSpacing = 16
        for i in range(len(series)):
            bc.bars[i].fillColor = colors.HexColor(paleta[i % len(paleta)])
        d.add(bc)

        leyenda = Legend()
        leyenda.x = 390
        leyenda.y = 110
        leyenda.dx = 8
        leyenda.dy = 8
        leyenda.fontSize = 7
        leyenda.alignment = "left"
        leyenda.colorNamePairs = [
            (colors.HexColor(paleta[i % len(paleta)]), nombre) for i, nombre in enumerate(series.keys())
        ]
        d.add(leyenda)
        return d

    @staticmethod
    def _grafico_pie(datos: Dict[str, float], titulo: str, colores_lista: Optional[List[str]] = None) -> Drawing:
        paleta = colores_lista or ["#1e3a8a", "#3b82f6", "#f97316", "#dc2626", "#16a34a", "#94a3b8"]
        d = Drawing(460, 180)
        d.add(String(0, 168, titulo, fontSize=8, fillColor=colors.HexColor("#334155")))
        pie = Pie()
        pie.x = 60
        pie.y = 10
        pie.width = 130
        pie.height = 130
        pie.data = list(datos.values())
        pie.labels = [f"{k} ({v})" for k, v in datos.items()]
        pie.slices.strokeWidth = 0.75
        pie.slices.strokeColor = colors.white
        pie.simpleLabels = 0
        pie.slices.fontSize = 7
        for i in range(len(datos)):
            pie.slices[i].fillColor = colors.HexColor(paleta[i % len(paleta)])
        d.add(pie)
        return d

    @staticmethod
    def _crear_pdf_base(
        titulo: str,
        columnas: list,
        filas: list,
        kpis: Optional[List[Tuple[str, str, str]]] = None,
        graficos: Optional[List[Drawing]] = None,
    ) -> io.BytesIO:
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=letter,
            rightMargin=24,
            leftMargin=24,
            topMargin=24,
            bottomMargin=24,
        )
        elements = []
        styles = getSampleStyleSheet()

        cell_style = ParagraphStyle(
            "CellStyle",
            parent=styles["Normal"],
            fontSize=7.5,
            leading=9,
            textColor=colors.HexColor("#0f172a"),
        )
        header_style = ParagraphStyle(
            "HeaderStyle",
            parent=styles["Normal"],
            fontSize=8,
            leading=10,
            textColor=colors.whitesmoke,
            fontName="Helvetica-Bold",
        )

        elements.append(Paragraph("<b>EL PRÍNCIPE</b>", styles["Heading1"]))
        elements.append(Paragraph(f"<b>Reporte:</b> {titulo}", styles["Heading2"]))
        elements.append(
            Paragraph(
                f"Fecha de emisión: {datetime.now().strftime('%d/%m/%Y %H:%M')}",
                styles["Normal"],
            )
        )
        elements.append(Spacer(1, 12))

        if kpis:
            elements.append(ReportService._kpi_cards(kpis))
            elements.append(Spacer(1, 16))

        if graficos:
            for grafico in graficos:
                elements.append(grafico)
                elements.append(Spacer(1, 8))
            elements.append(Spacer(1, 8))

        header_row = [Paragraph(str(col), header_style) for col in columnas]
        table_rows = [header_row]

        for fila in filas:
            table_rows.append(
                [
                    Paragraph(str(celda if celda is not None else "—"), cell_style)
                    for celda in fila
                ]
            )

        t = Table(table_rows, repeatRows=1)
        t.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E3A8A")),
                    ("BOTTOMPADDING", (0, 0), (-1, 0), 5),
                    ("TOPPADDING", (0, 0), (-1, 0), 5),
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    (
                        "ROWBACKGROUNDS",
                        (0, 1),
                        (-1, -1),
                        [colors.white, colors.HexColor("#F8FAFC")],
                    ),
                ]
            )
        )
        elements.append(t)

        doc.build(elements)
        buffer.seek(0)
        return buffer

    @staticmethod
    def _agregar_hoja_resumen_excel(
        writer,
        headers: List[str],
        filas: List[tuple],
        titulo: str,
        tipo: str = "bar",
        nombre_hoja: str = "Resumen",
    ) -> None:
        wb = writer.book
        ws = wb.create_sheet(nombre_hoja)
        ws.append(headers)
        for celda in ws[1]:
            celda.font = Font(bold=True)
        for fila in filas:
            ws.append(list(fila))

        n_filas = len(filas)
        n_series = len(headers) - 1
        data_ref = Reference(ws, min_col=2, max_col=1 + n_series, min_row=1, max_row=1 + n_filas)
        cat_ref = Reference(ws, min_col=1, min_row=2, max_row=1 + n_filas)

        chart = PieChart() if tipo == "pie" else BarChart()
        if tipo != "pie":
            chart.type = "col"
            chart.style = 10
        chart.title = titulo
        chart.add_data(data_ref, titles_from_data=True)
        chart.set_categories(cat_ref)
        chart.height = 8.5
        chart.width = 16
        ws.add_chart(chart, "E2")

    @staticmethod
    def generate_orders_excel(db: Session, desde=None, hasta=None) -> io.BytesIO:
        stmt = select(Pedido).options(selectinload(Pedido.cliente))
        if desde is not None:
            stmt = stmt.where(Pedido.fecha_pedido >= desde)
        if hasta is not None:
            stmt = stmt.where(Pedido.fecha_pedido <= hasta)
        orders = db.scalars(stmt).all()
        data = [
            {
                "Código": o.codigo_pedido,
                "Cliente": (
                    o.cliente.razon_social if o.cliente else f"ID {o.cliente_id}"
                ),
                "RUC/DNI": o.cliente.ruc_dni if o.cliente else "—",
                "Fecha": (
                    o.fecha_pedido.strftime("%Y-%m-%d %H:%M") if o.fecha_pedido else "—"
                ),
                "Fecha Entrega": (
                    o.fecha_entrega.strftime("%Y-%m-%d %H:%M")
                    if o.fecha_entrega
                    else "—"
                ),
                "Forma de Pago": str(o.forma_pago),
                "Estado": str(o.estado),
                "Monto Total (S/)": float(o.monto_total),
                "Observaciones": o.observaciones or "",
            }
            for o in orders
        ]
        df = pd.DataFrame(data)
        output = io.BytesIO()
        with pd.ExcelWriter(output, engine="openpyxl") as writer:
            df.to_excel(writer, index=False, sheet_name="Pedidos")

            orden_estados = ["Pendiente", "Aprobado", "Entregado", "Cancelado"]
            conteo = Counter(str(o.estado) for o in orders)
            ReportService._agregar_hoja_resumen_excel(
                writer,
                headers=["Estado", "Cantidad de Pedidos"],
                filas=[(e, conteo.get(e, 0)) for e in orden_estados],
                titulo="Pedidos por Estado",
                tipo="bar",
            )
        output.seek(0)
        return output

    @staticmethod
    def generate_inventory_excel(db: Session) -> io.BytesIO:
        products = db.scalars(select(Producto)).all()
        data = [
            {
                "SKU": p.sku,
                "Nombre": p.nombre,
                "Unidad": p.unidad_medida,
                "Stock Actual": p.stock_actual,
                "Stock Mínimo": p.stock_minimo,
                "Precio Costo (S/)": float(p.precio_costo),
                "Precio Venta (S/)": float(p.precio_unitario),
                "Rotación": str(p.nivel_rotacion),
                "Estado Stock": "Crítico" if p.stock_actual <= p.stock_minimo else "OK",
            }
            for p in products
        ]
        df = pd.DataFrame(data)
        output = io.BytesIO()
        with pd.ExcelWriter(output, engine="openpyxl") as writer:
            df.to_excel(writer, index=False, sheet_name="Inventario")

            conteo_rotacion = Counter(str(p.nivel_rotacion or "Sin definir") for p in products)
            ReportService._agregar_hoja_resumen_excel(
                writer,
                headers=["Nivel de Rotación", "Cantidad de Productos"],
                filas=list(conteo_rotacion.items()),
                titulo="Productos por Nivel de Rotación",
                tipo="pie",
            )
        output.seek(0)
        return output

    @staticmethod
    def generate_customers_excel(db: Session) -> io.BytesIO:
        customers = db.scalars(select(Cliente)).all()
        data = [
            {
                "Tipo Doc": "RUC" if c.ruc_dni and len(c.ruc_dni) == 11 else "DNI",
                "N° Documento": c.ruc_dni or "—",
                "Razón Social / Nombre": c.razon_social or "—",
                "Tipo Cliente": str(getattr(c, "tipo_cliente", "") or "—"),
                "Dirección": str(getattr(c, "direccion", "") or ""),
                "Distrito": str(getattr(c, "distrito", "") or ""),
                "Teléfono": str(getattr(c, "telefono", "") or ""),
                "Correo": str(getattr(c, "correo", "") or ""),
                "Estado": str(getattr(c, "estado", "") or "Activo"),
            }
            for c in customers
        ]
        df = pd.DataFrame(
            data
            if data
            else [
                {
                    "Tipo Doc": "—",
                    "N° Documento": "—",
                    "Razón Social / Nombre": "—",
                    "Tipo Cliente": "—",
                    "Dirección": "",
                    "Distrito": "",
                    "Teléfono": "",
                    "Correo": "",
                    "Estado": "—",
                }
            ]
        )
        output = io.BytesIO()
        with pd.ExcelWriter(output, engine="openpyxl") as writer:
            df.to_excel(writer, index=False, sheet_name="Clientes")

            if customers:
                conteo_tipo = Counter(str(getattr(c, "tipo_cliente", None) or "Sin definir") for c in customers)
                ReportService._agregar_hoja_resumen_excel(
                    writer,
                    headers=["Tipo de Cliente", "Cantidad"],
                    filas=list(conteo_tipo.items()),
                    titulo="Clientes por Tipo",
                    tipo="pie",
                )
        output.seek(0)
        return output

    @staticmethod
    def _datos_historial_cliente(db: Session, cliente_id: int):
        cliente = db.get(Cliente, cliente_id)
        if cliente is None:
            return None
        pedidos = db.scalars(
            select(Pedido)
            .options(selectinload(Pedido.detalles).selectinload(DetallePedido.producto))
            .where(Pedido.cliente_id == cliente_id)
            .order_by(Pedido.fecha_pedido.desc())
        ).all()
        solicitudes = db.scalars(
            select(SolicitudCliente)
            .where(SolicitudCliente.cliente_id == cliente_id)
            .order_by(SolicitudCliente.fecha_solicitud.desc())
        ).all()
        confirmados = [p for p in pedidos if p.estado in ("Aprobado", "Entregado")]
        total = sum(float(p.monto_total or 0) for p in confirmados)
        productos: Dict[str, List[float]] = {}
        for p in confirmados:
            for det in p.detalles:
                nombre = det.producto.nombre if det.producto else f"Producto #{det.producto_id}"
                acumulado = productos.setdefault(nombre, [0.0, 0.0])
                acumulado[0] += float(det.cantidad or 0)
                acumulado[1] += float(det.subtotal or 0)
        top_productos = sorted(productos.items(), key=lambda kv: -kv[1][1])
        fechas = [p.fecha_pedido.astimezone(LIMA_TZ) for p in confirmados]
        return {
            "cliente": cliente,
            "pedidos": pedidos,
            "solicitudes": solicitudes,
            "confirmados": confirmados,
            "total": total,
            "ticket": total / len(confirmados) if confirmados else 0.0,
            "productos": top_productos,
            "primera": min(fechas) if fechas else None,
            "ultima": max(fechas) if fechas else None,
        }

    @staticmethod
    def generate_customer_history_excel(db: Session, cliente_id: int) -> Optional[io.BytesIO]:
        d = ReportService._datos_historial_cliente(db, cliente_id)
        if d is None:
            return None
        c = d["cliente"]
        resumen = [
            ("Cliente", c.razon_social),
            ("RUC/DNI", c.ruc_dni or "—"),
            ("Tipo de cliente", str(c.tipo_cliente)),
            ("Dirección", c.direccion or "—"),
            ("Teléfono", c.telefono or "—"),
            ("Pedidos registrados", len(d["pedidos"])),
            ("Pedidos confirmados", len(d["confirmados"])),
            ("Total comprado (S/)", round(d["total"], 2)),
            ("Ticket promedio (S/)", round(d["ticket"], 2)),
            ("Primera compra", d["primera"].strftime("%d/%m/%Y") if d["primera"] else "—"),
            ("Última compra", d["ultima"].strftime("%d/%m/%Y") if d["ultima"] else "—"),
            ("Cotizaciones", len(d["solicitudes"])),
        ]
        pedidos = [
            {
                "Código": p.codigo_pedido,
                "Fecha": p.fecha_pedido.astimezone(LIMA_TZ).strftime("%d/%m/%Y %H:%M"),
                "Estado": p.estado,
                "Forma de pago": p.forma_pago,
                "Ítems": len(p.detalles),
                "Total (S/)": float(p.monto_total or 0),
            }
            for p in d["pedidos"]
        ]
        productos = [
            {"Producto": nombre, "Unidades": round(v[0], 2), "Monto (S/)": round(v[1], 2)}
            for nombre, v in d["productos"]
        ]
        cotizaciones = [
            {
                "Código": s.codigo_solicitud,
                "Fecha": s.fecha_solicitud.astimezone(LIMA_TZ).strftime("%d/%m/%Y %H:%M"),
                "Estado": s.estado,
            }
            for s in d["solicitudes"]
        ]
        output = io.BytesIO()
        with pd.ExcelWriter(output, engine="openpyxl") as writer:
            pd.DataFrame(resumen, columns=["Dato", "Valor"]).to_excel(writer, index=False, sheet_name="Resumen")
            pd.DataFrame(pedidos or [{"Código": "—"}]).to_excel(writer, index=False, sheet_name="Pedidos")
            pd.DataFrame(productos or [{"Producto": "—"}]).to_excel(writer, index=False, sheet_name="Productos")
            pd.DataFrame(cotizaciones or [{"Código": "—"}]).to_excel(writer, index=False, sheet_name="Cotizaciones")
            for hoja in writer.sheets.values():
                for columna in hoja.columns:
                    ancho = max(len(str(celda.value or "")) for celda in columna) + 2
                    hoja.column_dimensions[columna[0].column_letter].width = min(ancho, 48)
        output.seek(0)
        return output

    @staticmethod
    def generate_customer_history_pdf(db: Session, cliente_id: int) -> Optional[io.BytesIO]:
        d = ReportService._datos_historial_cliente(db, cliente_id)
        if d is None:
            return None
        c = d["cliente"]
        columnas = ["Código", "Fecha", "Estado", "Forma de pago", "Ítems", "Total S/"]
        filas = [
            [
                p.codigo_pedido,
                p.fecha_pedido.astimezone(LIMA_TZ).strftime("%d/%m/%Y"),
                p.estado,
                p.forma_pago,
                len(p.detalles),
                f"{float(p.monto_total or 0):,.2f}",
            ]
            for p in d["pedidos"]
        ] or [["—", "Sin pedidos", "—", "—", "—", "—"]]
        kpis = [
            ("Pedidos confirmados", str(len(d["confirmados"])), "#1E3A8A"),
            ("Total comprado", f"S/ {d['total']:,.2f}", "#15803d"),
            ("Ticket promedio", f"S/ {d['ticket']:,.2f}", "#7e22ce"),
            ("Última compra", d["ultima"].strftime("%d/%m/%Y") if d["ultima"] else "—", "#c2410c"),
        ]
        graficos = []
        if d["productos"]:
            top = {(n[:13] + "." if len(n) > 14 else n): round(v[1], 2) for n, v in d["productos"][:5]}
            graficos.append(ReportService._grafico_barras(top, "Productos más comprados (S/)"))
        return ReportService._crear_pdf_base(
            f"Historial de {c.razon_social}", columnas, filas, kpis=kpis, graficos=graficos
        )

    @staticmethod
    def generate_order_errors_excel(db: Session, desde=None, hasta=None) -> io.BytesIO:
        stmt = (
            select(DetallePedido)
            .join(Pedido)
            .options(
                selectinload(DetallePedido.pedido).selectinload(Pedido.cliente),
                selectinload(DetallePedido.producto),
            )
            .where(DetallePedido.tiene_error == True)
        )
        if desde is not None:
            stmt = stmt.where(Pedido.fecha_pedido >= desde)
        if hasta is not None:
            stmt = stmt.where(Pedido.fecha_pedido <= hasta)
        detalles_error = db.scalars(stmt).all()
        data = [
            {
                "N° Pedido": d.pedido.codigo_pedido if d.pedido else "—",
                "Cliente": (
                    d.pedido.cliente.razon_social
                    if d.pedido and d.pedido.cliente
                    else "—"
                ),
                "Producto": d.producto.nombre if d.producto else "—",
                "Cantidad": d.cantidad,
                "Tipo de Error": str(d.tipo_error),
                "Detalle / Observación": d.descripcion_error
                or d.pedido.observaciones
                or "",
                "Fecha Pedido": (
                    d.pedido.fecha_pedido.strftime("%Y-%m-%d")
                    if d.pedido and d.pedido.fecha_pedido
                    else "—"
                ),
                "Estado Pedido": str(d.pedido.estado) if d.pedido else "—",
            }
            for d in detalles_error
        ]
        df = pd.DataFrame(data)
        output = io.BytesIO()
        with pd.ExcelWriter(output, engine="openpyxl") as writer:
            df.to_excel(writer, index=False, sheet_name="Errores_Pedidos")

            if detalles_error:
                conteo_tipo_error = Counter(str(d.tipo_error) for d in detalles_error)
                ReportService._agregar_hoja_resumen_excel(
                    writer,
                    headers=["Tipo de Error", "Cantidad"],
                    filas=list(conteo_tipo_error.items()),
                    titulo="Distribución por Tipo de Error",
                    tipo="pie",
                )
        output.seek(0)
        return output

    @staticmethod
    def generate_actividad_comercial_excel(db: Session, sigla: str, dias: Optional[int] = None) -> io.BytesIO:
        from app.services.activity_indicator_service import ActivityIndicatorService

        campo_por_sigla = {
            "NCCA": "solicitudes",
            "NPP": "pedidos",
            "TPTD": "tiempo_promedio_decision_minutos",
        }
        columna_por_sigla = {
            "NCCA": "N° Cotizaciones de Clientes Atendidas (CA)",
            "NPP": "N° Pedidos Procesados (PA)",
            "TPTD": "Tiempo Promedio de Toma de Decisión (min)",
        }
        hoja_por_sigla = {"NCCA": "NCCA_por_dia", "NPP": "NPP_por_dia", "TPTD": "TPTD_por_dia"}

        campo = campo_por_sigla.get(sigla, "solicitudes")
        columna = columna_por_sigla.get(sigla, "Valor")

        serie = ActivityIndicatorService.get_serie_dias(db, cantidad_dias=dias)

        if sigla == "TPTD":
            return ReportService._tptd_por_dia_excel(db, serie)

        data = [
            {
                "N° Registro": f"Día {i + 1}",
                "Fecha": punto["label"],
                columna: (
                    round(punto[campo], 2) if punto[campo] is not None else "—"
                ),
            }
            for i, punto in enumerate(serie)
        ]
        df = pd.DataFrame(data if data else [{"N° Registro": "—", "Fecha": "—", columna: "—"}])

        output = io.BytesIO()
        with pd.ExcelWriter(output, engine="openpyxl") as writer:
            df.to_excel(writer, index=False, sheet_name=hoja_por_sigla.get(sigla, "Indicador"))
        output.seek(0)
        return output

    @staticmethod
    def _tptd_por_dia_excel(db: Session, serie: list) -> io.BytesIO:
        verde = PatternFill("solid", fgColor="E2EFDA")
        lado = Side(style="thin", color="000000")
        borde = Border(left=lado, right=lado, top=lado, bottom=lado)
        centro = Alignment(horizontal="center", vertical="center", wrap_text=True)
        izquierda = Alignment(horizontal="left", vertical="center", wrap_text=True)

        wb = Workbook()
        hoja = wb.active
        hoja.title = "Ficha_TPTD"
        for col, ancho in zip("ABCDE", (16, 16, 16, 16, 26)):
            hoja.column_dimensions[col].width = ancho

        def celda(ref, valor, relleno=None, negrita=False, alineacion=centro):
            c = hoja[ref]
            c.value = valor
            c.font = Font(bold=negrita)
            c.alignment = alineacion
            if relleno:
                c.fill = relleno
            return c

        def combinar(rango, valor, relleno=None, negrita=False, alineacion=centro):
            hoja.merge_cells(rango)
            celda(rango.split(":")[0], valor, relleno, negrita, alineacion)
            for fila in hoja[rango]:
                for c in fila:
                    c.border = borde
                    if relleno:
                        c.fill = relleno

        combinar("A1:E1", "FICHA DE OBSERVACIÓN", verde, True)
        combinar("A2:B2", "Indicador", verde, True)
        combinar("C2:E2", "Tiempo promedio de toma de decisiones", None, False, izquierda)
        combinar("A3:B3", "Investigador", verde, True)
        combinar("C3:E3", "Victor Jesus Nontol Nureña", None, False, izquierda)
        combinar("A4:B4", "Empresa", verde, True)
        combinar("C4:E4", "Productos El Príncipe", None, False, izquierda)

        celda("A5", "Variable", verde, True)
        celda("B5", "Medida", verde, True)
        combinar("C5:E5", "Fórmula del indicador", verde, True)
        celda("A6", "Toma de decisiones")
        celda("B6", "Segundos")
        combinar(
            "C6:E6",
            "TPTD = Σ (Hora de fin − Hora de inicio) / N.° de decisiones tomadas",
        )
        hoja.row_dimensions[6].height = 34

        for col, titulo in zip(
            "ABCDE", ("Ítem", "Fecha", "Hora de Inicio", "Hora de Fin", "Tiempo de toma de decisión")
        ):
            celda(f"{col}7", titulo, verde, True)
        hoja.row_dimensions[7].height = 32

        fila = 8
        suma_ponderada = 0.0
        total_pedidos = 0
        for i, punto in enumerate(serie):
            valor = punto["tiempo_promedio_decision_segundos"]
            n = punto.get("pedidos_con_decision", 0)
            if valor is not None:
                suma_ponderada += valor * n
                total_pedidos += n
            celda(f"A{fila}", i + 1, negrita=True)
            celda(f"B{fila}", punto.get("fecha_completa") or punto["label"], negrita=True)
            celda(f"C{fila}", punto.get("hora_inicio_decision") or "—")
            celda(f"D{fila}", punto.get("hora_fin_decision") or "—")
            celda(f"E{fila}", round(valor) if valor is not None else "Sin pedidos")
            fila += 1

        if total_pedidos:
            combinar(
                f"A{fila}:D{fila}",
                f"TPTD = Σ / N  (N = {total_pedidos} decisiones tomadas)",
                verde,
                True,
            )
            celda(f"E{fila}", round(suma_ponderada / total_pedidos), verde, True)
            fila += 1

        for fila_celdas in hoja[f"A1:E{fila - 1}"]:
            for c in fila_celdas:
                c.border = borde

        ReportService._hoja_detalle_tptd(db, wb, serie, verde, borde, centro)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output

    @staticmethod
    def _hoja_detalle_tptd(db: Session, wb, serie: list, verde, borde, centro) -> None:
        hoja = wb.create_sheet("Detalle_por_pedido")
        for col, ancho in zip("ABCDE", (8, 16, 22, 22, 26)):
            hoja.column_dimensions[col].width = ancho

        titulos = ("Ítem", "Fecha", "Hora de Inicio", "Hora de Fin", "Tiempo de toma de decisión (seg)")
        for col, titulo in zip("ABCDE", titulos):
            c = hoja[f"{col}1"]
            c.value = titulo
            c.font = Font(bold=True)
            c.fill = verde
            c.alignment = centro
            c.border = borde
        hoja.row_dimensions[1].height = 32

        filas = []
        if serie and serie[0].get("fecha_completa"):
            primer_dia = datetime.strptime(serie[0]["fecha_completa"], "%d/%m/%Y").replace(tzinfo=LIMA_TZ)
            stmt = (
                select(Pedido.fecha_aprobacion, SolicitudCliente.hora_apertura_modal, SolicitudCliente.fecha_solicitud)
                .join(SolicitudCliente, SolicitudCliente.id == Pedido.solicitud_id)
                .where(Pedido.fecha_aprobacion.isnot(None), Pedido.fecha_aprobacion >= primer_dia)
                .order_by(Pedido.fecha_aprobacion.asc())
            )
            filas = db.execute(stmt).all()

        for i, (aprobacion, apertura, fecha_solicitud) in enumerate(filas):
            fin = aprobacion.astimezone(LIMA_TZ)
            inicio = (apertura or fecha_solicitud).astimezone(LIMA_TZ)
            segundos = round((fin - inicio).total_seconds())
            valores = (
                i + 1,
                fin.strftime("%d/%m/%Y"),
                inicio.strftime("%d/%m/%Y %H:%M:%S"),
                fin.strftime("%d/%m/%Y %H:%M:%S"),
                segundos,
            )
            for col, valor in zip("ABCDE", valores):
                c = hoja[f"{col}{i + 2}"]
                c.value = valor
                c.alignment = centro
                c.border = borde

    @staticmethod
    def generate_tpd_por_pedido_excel(db: Session) -> io.BytesIO:
        stmt = (
            select(Pedido, SolicitudCliente)
            .join(SolicitudCliente, SolicitudCliente.id == Pedido.solicitud_id)
            .where(Pedido.fecha_aprobacion.isnot(None))
            .order_by(SolicitudCliente.fecha_solicitud.asc())
        )
        filas = db.execute(stmt).all()

        data = []
        for i, (pedido, solicitud) in enumerate(filas):
            inicio = (solicitud.hora_apertura_modal or solicitud.fecha_solicitud).astimezone(LIMA_TZ)
            fin = pedido.fecha_aprobacion.astimezone(LIMA_TZ)
            segundos = round((fin - inicio).total_seconds())
            apertura_pedido = pedido.hora_apertura_modal.astimezone(LIMA_TZ) if pedido.hora_apertura_modal else None
            fin_solicitud = solicitud.fecha_solicitud.astimezone(LIMA_TZ)
            seg_solicitud = round((fin_solicitud - inicio).total_seconds())
            seg_pedido = round((fin - apertura_pedido).total_seconds()) if apertura_pedido else None
            data.append(
                {
                    "Ítem": i + 1,
                    "Hora de Inicio (Apertura del registro de cotización)": inicio.strftime("%d/%m %H:%M"),
                    "Hora de Apertura (Nuevo pedido)": apertura_pedido.strftime("%d/%m %H:%M") if apertura_pedido else "—",
                    "Hora de Fin (Aprobación del pedido)": fin.strftime("%d/%m %H:%M"),
                    "Etapa Cotización (seg)": seg_solicitud,
                    "Etapa Pedido (seg)": seg_pedido if seg_pedido is not None else "—",
                    "Tiempo de Toma de Decisión (seg)": segundos,
                }
            )

        df = pd.DataFrame(
            data
            if data
            else [
                {
                    "Ítem": "—",
                    "Hora de Inicio (Apertura del registro de cotización)": "—",
                    "Hora de Apertura (Nuevo pedido)": "—",
                    "Hora de Fin (Aprobación del pedido)": "—",
                    "Etapa Cotización (seg)": "—",
                    "Etapa Pedido (seg)": "—",
                    "Tiempo de Toma de Decisión (seg)": "—",
                }
            ]
        )

        output = io.BytesIO()
        with pd.ExcelWriter(output, engine="openpyxl") as writer:
            df.to_excel(writer, index=False, sheet_name="TPD_por_pedido")
        output.seek(0)
        return output

    @staticmethod
    def generate_indicators_excel(db: Session) -> io.BytesIO:
        pedidos = list(
            db.scalars(select(Pedido).options(selectinload(Pedido.detalles))).all()
        )
        total_pedidos_items = sum(len(p.detalles) for p in pedidos)
        detalles_con_error = sum(
            1 for p in pedidos for d in p.detalles if d.tiene_error
        )
        nepp = (
            (detalles_con_error / total_pedidos_items)
            if total_pedidos_items > 0
            else 0.0
        )

        indicador = IndicatorService.get_latest(db) or IndicatorService.calculate_and_save(
            db, resumen="Cálculo automático para reporte Excel"
        )
        pfcc = float(indicador.valor_pfcc or 0)
        ntdc = float(indicador.valor_ntdc or 0)

        data = [
            {
                "Sigla": "NEPP",
                "Indicador": "Número de errores en los productos pedidos",
                "Fórmula": "NEPP = TEPP ÷ TPP",
                "Valor Numérico": round(nepp, 4),
                "Formato Presentación": f"{round(nepp * 100, 2)}%",
                "Meta": "< 0.05",
                "Detalle Muestra": f"{detalles_con_error} errores de {total_pedidos_items} ítems registrados",
            },
            {
                "Sigla": "PFCC",
                "Indicador": "Porcentaje de fallas al definir condiciones comerciales",
                "Fórmula": "PFCC = (TCCF ÷ TCCD) × 100",
                "Valor Numérico": round(pfcc / 100, 4),
                "Formato Presentación": f"{round(pfcc, 2)}%",
                "Meta": "< 10%",
                "Detalle Muestra": f"{indicador.total_fallas_condiciones} fallas de {indicador.total_condiciones_pactadas} condiciones registradas",
            },
            {
                "Sigla": "NTDC",
                "Indicador": "Nivel de toma de decisiones comerciales",
                "Fórmula": "NTDC = (TDCE ÷ TDCT) × 100",
                "Valor Numérico": round(ntdc / 100, 4),
                "Formato Presentación": f"{round(ntdc, 2)}%",
                "Meta": "≥ 75%",
                "Detalle Muestra": (
                    f"{indicador.total_decisiones_efectivas} decisiones efectivas de "
                    f"{indicador.total_decisiones_evaluadas} registradas "
                    f"({indicador.total_decisiones_corregidas} corregidas tras incidencia)"
                ),
            },
        ]
        df = pd.DataFrame(data)
        output = io.BytesIO()
        with pd.ExcelWriter(output, engine="openpyxl") as writer:
            df.to_excel(writer, index=False, sheet_name="Indicadores")

            ReportService._agregar_hoja_resumen_excel(
                writer,
                headers=["Indicador", "Actual (%)", "Meta (%)"],
                filas=[
                    ("NEPP", round(nepp * 100, 2), 5),
                    ("PFCC", round(pfcc, 2), 10),
                    ("NTDC", round(ntdc, 2), 75),
                ],
                titulo="Indicadores vs. Meta (%)",
                tipo="bar",
            )
        output.seek(0)
        return output

    @staticmethod
    def generate_orders_pdf(db: Session, desde=None, hasta=None) -> io.BytesIO:
        stmt = select(Pedido).options(selectinload(Pedido.cliente))
        if desde is not None:
            stmt = stmt.where(Pedido.fecha_pedido >= desde)
        if hasta is not None:
            stmt = stmt.where(Pedido.fecha_pedido <= hasta)
        orders = db.scalars(stmt).all()
        columnas = ["Código", "Cliente", "Fecha", "Forma Pago", "Estado", "Total (S/)"]
        filas = [
            [
                str(o.codigo_pedido),
                str(o.cliente.razon_social if o.cliente else f"ID {o.cliente_id}"),
                o.fecha_pedido.strftime("%Y-%m-%d") if o.fecha_pedido else "—",
                str(o.forma_pago),
                str(o.estado),
                f"S/ {float(o.monto_total):,.2f}",
            ]
            for o in orders
        ]

        total_pedidos = len(orders)
        monto_total = sum(float(o.monto_total or 0) for o in orders)
        ticket_prom = monto_total / total_pedidos if total_pedidos else 0
        kpis = [
            ("Total Pedidos", str(total_pedidos), "#1E3A8A"),
            ("Monto Total", f"S/ {monto_total:,.2f}", "#16a34a"),
            ("Ticket Promedio", f"S/ {ticket_prom:,.2f}", "#ea580c"),
        ]

        orden_estados = ["Pendiente", "Aprobado", "Entregado", "Cancelado"]
        conteo = Counter(str(o.estado) for o in orders)
        grafico = ReportService._grafico_barras(
            {e: conteo.get(e, 0) for e in orden_estados}, "Pedidos por Estado"
        )

        return ReportService._crear_pdf_base(
            "Reporte de Pedidos", columnas, filas, kpis=kpis, graficos=[grafico]
        )

    @staticmethod
    def generate_inventory_pdf(db: Session) -> io.BytesIO:
        products = db.scalars(select(Producto)).all()
        columnas = [
            "SKU",
            "Producto",
            "Unidad",
            "Stock",
            "Mínimo",
            "Precio (S/)",
            "Rotación",
        ]
        filas = [
            [
                str(p.sku),
                str(p.nombre),
                str(p.unidad_medida),
                str(p.stock_actual),
                str(p.stock_minimo),
                f"S/ {float(p.precio_unitario):.2f}",
                str(p.nivel_rotacion),
            ]
            for p in products
        ]

        total_productos = len(products)
        criticos = sum(1 for p in products if (p.stock_actual or 0) <= (p.stock_minimo or 0))
        valor_inventario = sum(float(p.precio_costo or 0) * (p.stock_actual or 0) for p in products)
        kpis = [
            ("Total Productos", str(total_productos), "#1E3A8A"),
            ("Stock Crítico", str(criticos), "#dc2626"),
            ("Valor Inventario", f"S/ {valor_inventario:,.2f}", "#16a34a"),
        ]

        conteo_rotacion = Counter(str(p.nivel_rotacion or "Sin definir") for p in products)
        graficos = [ReportService._grafico_pie(dict(conteo_rotacion), "Productos por Nivel de Rotación")] if products else []

        return ReportService._crear_pdf_base(
            "Catálogo de Inventario", columnas, filas, kpis=kpis, graficos=graficos
        )

    @staticmethod
    def generate_customers_pdf(db: Session) -> io.BytesIO:
        customers = db.scalars(select(Cliente)).all()
        columnas = ["N° Doc", "Razón Social", "Tipo", "Distrito", "Teléfono", "Estado"]
        filas = [
            [
                str(c.ruc_dni or "—"),
                str(c.razon_social or "—"),
                str(getattr(c, "tipo_cliente", "") or "—"),
                str(getattr(c, "distrito", "") or "—"),
                str(getattr(c, "telefono", "") or "—"),
                str(getattr(c, "estado", "") or "Activo"),
            ]
            for c in customers
        ]
        if not filas:
            filas = [["—", "Sin registros", "—", "—", "—", "—"]]

        kpis = [("Total Clientes", str(len(customers)), "#1E3A8A")]
        graficos = []
        if customers:
            conteo_tipo = Counter(str(getattr(c, "tipo_cliente", None) or "Sin definir") for c in customers)
            graficos.append(ReportService._grafico_pie(dict(conteo_tipo), "Clientes por Tipo"))

        return ReportService._crear_pdf_base(
            "Cartera de Clientes", columnas, filas, kpis=kpis, graficos=graficos
        )

    @staticmethod
    def generate_order_errors_pdf(db: Session, desde=None, hasta=None) -> io.BytesIO:
        stmt = (
            select(DetallePedido)
            .join(Pedido)
            .options(
                selectinload(DetallePedido.pedido).selectinload(Pedido.cliente),
                selectinload(DetallePedido.producto),
            )
            .where(DetallePedido.tiene_error == True)
        )
        if desde is not None:
            stmt = stmt.where(Pedido.fecha_pedido >= desde)
        if hasta is not None:
            stmt = stmt.where(Pedido.fecha_pedido <= hasta)
        detalles_error = db.scalars(stmt).all()
        columnas = ["N° Pedido", "Cliente", "Producto", "Cant.", "Tipo Error", "Fecha"]
        filas = [
            [
                str(d.pedido.codigo_pedido if d.pedido else "—"),
                str(
                    d.pedido.cliente.razon_social
                    if d.pedido and d.pedido.cliente
                    else "—"
                ),
                str(d.producto.nombre if d.producto else "—"),
                str(d.cantidad),
                str(d.tipo_error or "—"),
                (
                    d.pedido.fecha_pedido.strftime("%Y-%m-%d")
                    if d.pedido and d.pedido.fecha_pedido
                    else "—"
                ),
            ]
            for d in detalles_error
        ]

        total_errores = len(detalles_error)
        pedidos_afectados = len({d.pedido_id for d in detalles_error})
        kpis = [
            ("Total Errores", str(total_errores), "#dc2626"),
            ("Pedidos Afectados", str(pedidos_afectados), "#ea580c"),
        ]
        graficos = []
        if detalles_error:
            conteo_tipo_error = Counter(str(d.tipo_error) for d in detalles_error)
            graficos.append(
                ReportService._grafico_pie(
                    dict(conteo_tipo_error),
                    "Distribución por Tipo de Error",
                    colores_lista=["#dc2626", "#ea580c", "#f59e0b", "#b91c1c"],
                )
            )

        return ReportService._crear_pdf_base(
            "Pedidos con Errores Detectados", columnas, filas, kpis=kpis, graficos=graficos
        )

    @staticmethod
    def generate_indicators_pdf(db: Session) -> io.BytesIO:
        pedidos = list(
            db.scalars(select(Pedido).options(selectinload(Pedido.detalles))).all()
        )
        total_pedidos_items = sum(len(p.detalles) for p in pedidos)
        detalles_con_error = sum(
            1 for p in pedidos for d in p.detalles if d.tiene_error
        )
        nepp = (
            (detalles_con_error / total_pedidos_items)
            if total_pedidos_items > 0
            else 0.0
        )

        indicador = IndicatorService.get_latest(db) or IndicatorService.calculate_and_save(
            db, resumen="Cálculo automático para reporte PDF"
        )
        pfcc = float(indicador.valor_pfcc or 0)
        ntdc = float(indicador.valor_ntdc or 0)

        columnas = ["Sigla", "Indicador", "Fórmula", "Valor", "Meta", "Muestra"]
        filas = [
            [
                "NEPP",
                "Errores en productos pedidos",
                "TEPP ÷ TPP",
                f"{round(nepp, 4)} ({round(nepp * 100, 2)}%)",
                "< 0.05",
                f"{detalles_con_error} / {total_pedidos_items}",
            ],
            [
                "PFCC",
                "Fallas en condiciones comerciales",
                "(TCCF ÷ TCCD) × 100",
                f"{round(pfcc, 2)}%",
                "< 10%",
                f"{indicador.total_fallas_condiciones} / {indicador.total_condiciones_pactadas}",
            ],
            [
                "NTDC",
                "Toma de decisiones comerciales",
                "(TDCE ÷ TDCT) × 100",
                f"{round(ntdc, 2)}%",
                "≥ 75%",
                f"{indicador.total_decisiones_efectivas} / {indicador.total_decisiones_evaluadas}",
            ],
        ]

        kpis = [
            ("NEPP", f"{round(nepp * 100, 2)}%", "#dc2626" if nepp > 0.10 else ("#ea580c" if nepp > 0.05 else "#16a34a")),
            ("PFCC", f"{round(pfcc, 2)}%", "#dc2626" if pfcc > 25 else ("#ea580c" if pfcc > 10 else "#16a34a")),
            ("NTDC", f"{round(ntdc, 2)}%", "#16a34a" if ntdc >= 75 else ("#ea580c" if ntdc >= 50 else "#dc2626")),
        ]
        grafico = ReportService._grafico_barras_comparativo(
            ["NEPP", "PFCC", "NTDC"],
            {
                "Actual (%)": [round(nepp * 100, 2), round(pfcc, 2), round(ntdc, 2)],
                "Meta (%)": [5, 10, 75],
            },
            "Indicadores vs. Meta (%)",
        )

        return ReportService._crear_pdf_base(
            "Indicadores Comerciales (NEPP · PFCC · NTDC)",
            columnas,
            filas,
            kpis=kpis,
            graficos=[grafico],
        )