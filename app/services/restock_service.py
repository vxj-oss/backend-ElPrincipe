import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.time import ahora_lima
from app.models.category import Categoria
from app.models.product import Producto

COLUMNAS = (
    ("N°", 6),
    ("SKU", 16),
    ("Producto", 44),
    ("Categoría", 22),
    ("Stock actual", 13),
    ("Stock mínimo", 13),
    ("Unidades a pedir", 17),
    ("Costo unit. (S/)", 16),
    ("Costo estimado (S/)", 19),
    ("Rotación", 11),
)


def generar_reposicion_excel(db: Session, factor: int = 2) -> io.BytesIO:
    filas = db.execute(
        select(Producto, Categoria.nombre)
        .outerjoin(Categoria, Producto.categoria_id == Categoria.id)
        .where(Producto.activo.is_(True), Producto.stock_actual <= Producto.stock_minimo)
        .order_by(Producto.stock_actual.asc(), Producto.nombre.asc())
    ).all()

    pendientes = []
    for producto, categoria in filas:
        unidades = max(0, factor * (producto.stock_minimo or 0) - (producto.stock_actual or 0))
        if unidades > 0:
            pendientes.append((producto, categoria or "Sin categoría", unidades))

    wb = Workbook()
    hoja = wb.active
    hoja.title = "Reposición"

    verde = PatternFill("solid", fgColor="E2EFDA")
    azul = PatternFill("solid", fgColor="1E3A8A")
    lado = Side(style="thin", color="CBD5E1")
    borde = Border(left=lado, right=lado, top=lado, bottom=lado)
    centro = Alignment(horizontal="center", vertical="center", wrap_text=True)

    hoja.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(COLUMNAS))
    hoja["A1"] = "LISTA DE REPOSICIÓN — EL PRÍNCIPE"
    hoja["A1"].font = Font(bold=True, size=14)
    hoja.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(COLUMNAS))
    hoja["A2"] = (
        f"Emitida el {ahora_lima().strftime('%d/%m/%Y %H:%M')} · "
        f"Se repone hasta {factor} veces el stock mínimo"
    )
    hoja["A2"].font = Font(color="64748B")

    for posicion, (titulo, ancho) in enumerate(COLUMNAS, start=1):
        celda = hoja.cell(row=4, column=posicion, value=titulo)
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = azul
        celda.alignment = centro
        celda.border = borde
        hoja.column_dimensions[celda.column_letter].width = ancho
    hoja.row_dimensions[4].height = 28

    if not pendientes:
        hoja.merge_cells(start_row=5, start_column=1, end_row=5, end_column=len(COLUMNAS))
        hoja["A5"] = "No hay productos bajo el stock mínimo."
        hoja["A5"].alignment = centro
    else:
        fila = 5
        total_unidades = 0
        total_costo = 0.0
        for numero, (producto, categoria, unidades) in enumerate(pendientes, start=1):
            costo = float(producto.precio_costo or 0)
            estimado = costo * unidades
            total_unidades += unidades
            total_costo += estimado
            valores = (
                numero,
                producto.sku,
                producto.nombre,
                categoria,
                producto.stock_actual or 0,
                producto.stock_minimo or 0,
                unidades,
                costo,
                estimado,
                str(producto.nivel_rotacion),
            )
            for posicion, valor in enumerate(valores, start=1):
                celda = hoja.cell(row=fila, column=posicion, value=valor)
                celda.border = borde
                celda.alignment = Alignment(
                    horizontal="left" if posicion in (3, 4) else "center", vertical="center"
                )
                if posicion in (8, 9):
                    celda.number_format = "#,##0.00"
            hoja.cell(row=fila, column=7).font = Font(bold=True)
            fila += 1

        hoja.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=6)
        hoja.cell(row=fila, column=1, value="TOTAL").font = Font(bold=True)
        hoja.cell(row=fila, column=1).alignment = Alignment(horizontal="right")
        hoja.cell(row=fila, column=7, value=total_unidades).font = Font(bold=True)
        hoja.cell(row=fila, column=9, value=total_costo).font = Font(bold=True)
        hoja.cell(row=fila, column=9).number_format = "#,##0.00"
        for posicion in range(1, len(COLUMNAS) + 1):
            celda = hoja.cell(row=fila, column=posicion)
            celda.fill = verde
            celda.border = borde
            if posicion in (7, 9):
                celda.alignment = Alignment(horizontal="center")

    hoja.freeze_panes = "A5"

    salida = io.BytesIO()
    wb.save(salida)
    salida.seek(0)
    return salida
